"""Excel sheet for Raid BiS comparisons."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from inventory_parser.excel_theme import (
    FILL_HEADER,
    FILL_LABEL,
    FILL_SHEET,
    FONT_BODY,
    FONT_HEADER,
    FONT_LEGEND,
    FONT_LINK,
    FONT_ON_STATUS,
    FONT_SECTION,
    STATUS_FILLS,
    evolver_fill,
)
from inventory_parser.items import EQRESOURCE_ITEM_URL
from inventory_parser.raid_bis.build import RaidBisExport
from inventory_parser.raid_bis.compare import format_stat_deltas
from inventory_parser.raid_bis.ore_needs import build_ore_needs_matrix
from inventory_parser.slot2_augs.html import format_catalog_fetched_at

SHEET_NAME = "Raid BiS"
MISSING_ORES_SHEET_NAME = "Missing Ores"

_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_HEADER = Alignment(horizontal="center", vertical="center")


def append_raid_bis_sheet(wb, bundle: RaidBisExport) -> None:
    ws = wb.create_sheet(SHEET_NAME)
    ws.sheet_properties.tabColor = "4A3520"
    headers = [
        "Character",
        "Class",
        "Slot",
        "Status",
        "Current",
        "Best in slot",
        "Tier",
        "Vendor cost",
        "Vendor item",
        "Stat changes",
        "Notes",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(1, col, h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")

    row = 2
    for ch in bundle.characters:
        for slot in ch.slots:
            fill = STATUS_FILLS.get(slot.status)
            ws.cell(row, 1, ch.display_name)
            ws.cell(row, 2, ch.class_abbr or "")
            ws.cell(row, 3, slot.gear_slot)
            status_cell = ws.cell(row, 4, slot.status)
            if fill:
                status_cell.fill = fill
                status_cell.font = FONT_ON_STATUS
            _item_cell(ws.cell(row, 5), slot.current_name, slot.current_id)
            _item_cell(ws.cell(row, 6), slot.recommended_name, slot.recommended_id)
            ws.cell(row, 7, slot.recommended_tier)
            if slot.vendor_cost is not None:
                ws.cell(row, 8, slot.vendor_cost)
            _item_cell(ws.cell(row, 9), slot.vendor_item_name, slot.vendor_item_id)
            ws.cell(row, 10, format_stat_deltas(slot.deltas, class_abbr=ch.class_abbr))
            ws.cell(row, 11, slot.note)
            row += 1

        if ch.total_deltas:
            ws.cell(row, 1, ch.display_name)
            ws.cell(row, 3, "TOTAL")
            ws.cell(row, 10, format_stat_deltas(ch.total_deltas, class_abbr=ch.class_abbr))
            row += 1

    widths = [22, 8, 12, 12, 42, 42, 8, 14, 36, 48, 36]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    fetched = format_catalog_fetched_at(bundle.catalog.fetched_at)
    cache_note = " (cache)" if bundle.catalog.from_cache else ""
    ws.cell(row + 1, 1, f"Catalog fetched {fetched}{cache_note}")


def append_missing_ores_sheet(wb, bundle: RaidBisExport) -> None:
    """Character × ore matrix of raid-vendor ores still needed for Raid BiS."""
    matrix = build_ore_needs_matrix(bundle)
    ws = wb.create_sheet(MISSING_ORES_SHEET_NAME)
    ws.sheet_properties.tabColor = "3A3350"
    n_chars = len(matrix.characters)
    summary_cols = max(2, 1 + n_chars)

    title = ws.cell(1, 1, MISSING_ORES_SHEET_NAME)
    title.font = FONT_SECTION
    title.fill = FILL_HEADER
    title.alignment = _ALIGN
    sub = ws.cell(
        2,
        1,
        "Raid vendor ores still needed for Raid BiS upgrades · "
        "purple = Evolver slot (excluded from Total)",
    )
    sub.font = FONT_LEGEND
    sub.fill = FILL_HEADER
    sub.alignment = _ALIGN
    for col in range(1, summary_cols + 1):
        ws.cell(1, col).fill = FILL_HEADER
        ws.cell(2, col).fill = FILL_HEADER
    if summary_cols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=summary_cols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=summary_cols)
    ws.row_dimensions[1].height = 24

    header_row = 4
    ore_header = ws.cell(header_row, 1, "Ore")
    ore_header.font = FONT_HEADER
    ore_header.fill = FILL_HEADER
    ore_header.alignment = _ALIGN_HEADER
    for col, name in enumerate(matrix.characters, start=2):
        cell = ws.cell(header_row, col, name)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = _ALIGN_HEADER

    purple = evolver_fill()
    row = header_row + 1
    for ore_row in matrix.rows:
        name_cell = ws.cell(row, 1, ore_row.name)
        name_cell.font = FONT_BODY
        name_cell.fill = FILL_LABEL
        name_cell.alignment = _ALIGN
        if ore_row.item_id > 0:
            name_cell.hyperlink = EQRESOURCE_ITEM_URL.format(item_id=ore_row.item_id)
            name_cell.font = FONT_LINK
        for col_i, count in enumerate(ore_row.counts):
            col = col_i + 2
            has_evolver = ore_row.evolvers[col_i]
            cell = ws.cell(row=row, column=col)
            cell.alignment = _ALIGN_CENTER
            if count > 0:
                cell.value = count
                cell.font = FONT_BODY
            if has_evolver:
                cell.fill = purple
                if count <= 0:
                    cell.value = "Evolver"
                    cell.font = Font(name="Calibri", size=11, color="E4E6EB")
            else:
                cell.fill = FILL_SHEET
        row += 1

    if matrix.rows:
        total_label = ws.cell(row, 1, "Total")
        total_label.font = FONT_HEADER
        total_label.fill = FILL_HEADER
        total_label.alignment = _ALIGN
        for col_i, total in enumerate(matrix.totals):
            cell = ws.cell(row, col_i + 2, total if total else "")
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = _ALIGN_CENTER
        row += 1

    note = ws.cell(
        row + 1,
        1,
        "Counts are ores for Raid BiS upgrades only (T2 linings/clasps/cloths). "
        "Already BiS and Evolver slots are excluded from counts and Total. Bag stock is not subtracted "
        "(see Unmade Gear). T1 finished vendor items and Diminished containers are omitted.",
    )
    note.font = FONT_LEGEND
    note.fill = FILL_SHEET
    note.alignment = _ALIGN
    if summary_cols > 1:
        ws.merge_cells(
            start_row=row + 1, start_column=1, end_row=row + 1, end_column=summary_cols
        )

    ws.column_dimensions["A"].width = 36.0
    for col in range(2, 2 + n_chars):
        label = matrix.characters[col - 2] if col - 2 < n_chars else ""
        ws.column_dimensions[get_column_letter(col)].width = max(
            11.0, min(22.0, len(label) + 2.0)
        )


def _item_cell(cell, name: str | None, item_id: int | None) -> None:
    if not name:
        cell.value = ""
        return
    cell.value = name
    if item_id and item_id > 0:
        cell.hyperlink = EQRESOURCE_ITEM_URL.format(item_id=item_id)
        cell.font = FONT_LINK
